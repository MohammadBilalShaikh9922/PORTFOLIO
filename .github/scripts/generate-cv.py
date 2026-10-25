
from pathlib import Path
from copy import deepcopy
import re
import subprocess

from bs4 import BeautifulSoup
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path.cwd()
TEMPLATE = ROOT / "Mohammad_Bilal_Shaikh.docx"
HTML_FILE = ROOT / "index.html"
OUTPUT_DOCX = ROOT / "Mohammad_Bilal_Shaikh_Generated.docx"
OUTPUT_PDF = ROOT / "portfolio.pdf"


def clean(element):
    return " ".join(element.stripped_strings) if element else ""


def extract_portfolio():
    soup = BeautifulSoup(HTML_FILE.read_text(encoding="utf-8"), "html.parser")

    contact = soup.select_one(".contact-links")
    email = soup.select_one('a[href^="mailto:"]')
    phone = contact.select_one("a[href^='tel:']") if contact else None

    if not phone and contact:
        for item in contact.select("div"):
            if item.find(string=re.compile("Phone", re.I)):
                phone = item.find("strong")
                break

    linkedin = soup.select_one('a[href*="linkedin.com/in/"]')
    github = soup.select_one('a[href*="github.com/"]')

    name = clean(soup.select_one("h1")).replace("  ", " ")

    contact_line = " | ".join(filter(None, [
        clean(phone),
        clean(email),
        linkedin.get("href", "").replace("https://", "").replace("www.", "")
        if linkedin else "",
        github.get("href", "").replace("https://", "").replace("www.", "")
        if github else "",
        "Pune, Maharashtra, India",
    ]))

    education = []
    for card in soup.select(".education-card"):
        education.append({
            "date": clean(card.select_one(".edu-year")),
            "degree": clean(card.select_one("h3")),
            "institution": clean(card.select_one("p")),
            "details": [
                clean(x) for x in card.select(".sgpa-row span")
            ] or [
                clean(x) for x in card.select("strong")
            ],
        })

    experience = []
    for item in soup.select(".timeline-item"):
        experience.append({
            "date": clean(item.select_one(".timeline-meta")),
            "role": clean(item.select_one("h3")),
            "company": clean(item.select_one(".company")),
            "bullets": [clean(x) for x in item.select("li")],
        })

    projects = []
    for card in soup.select(".project-card"):
        projects.append({
            "name": clean(card.select_one("h3")),
            "type": clean(card.select_one(".project-type")),
            "description": clean(card.select_one(":scope > p:not(.project-type)")),
            "tags": [clean(x) for x in card.select(".tags span")],
        })

    certifications = [
        clean(x) for x in soup.select(".cert-list span")
    ]

    skills = {}
    for block in soup.select(".skill-block"):
        heading = clean(block.select_one("h3"))
        skills[heading] = [
            clean(x) for x in block.select(".skill-pills span")
        ]

    languages = clean(soup.select_one(".language-card h2"))

    return {
        "name": name,
        "contact": contact_line,
        "education": education,
        "experience": experience,
        "projects": projects,
        "certifications": certifications,
        "skills": skills,
        "languages": languages,
    }


def paragraph_text(paragraph, value):
    """Replace text while retaining the paragraph's formatting."""
    first_run = paragraph.runs[0] if paragraph.runs else None
    run_properties = (
        deepcopy(first_run._r.rPr)
        if first_run is not None and first_run._r.rPr is not None
        else None
    )

    paragraph.clear()
    run = paragraph.add_run(value)

    if run_properties is not None:
        run._r.insert(0, run_properties)


def make_paragraph(document, prototype, value):
    """Copy a paragraph's original layout, style, and bullet formatting."""
    element = deepcopy(prototype._p)
    document._body._body.append(element)

    from docx.text.paragraph import Paragraph
    paragraph = Paragraph(element, document._body)
    paragraph_text(paragraph, value)
    return paragraph


def generate():
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"CV template missing: {TEMPLATE}")
    if not HTML_FILE.exists():
        raise FileNotFoundError(f"Portfolio HTML missing: {HTML_FILE}")

    data = extract_portfolio()
    document = Document(TEMPLATE)
    original = list(document.paragraphs)

    def find(text):
        for paragraph in original:
            if paragraph.text.strip() == text:
                return paragraph
        raise ValueError(f"Template paragraph not found: {text}")

    # Reuse the supplied Word document's original paragraph formatting.
    heading = find("EDUCATION")
    education_title = find(
        "Marathwada Mitramandal Institute of TechnologyAug. 2023 – May 2027"
    )
    education_detail = find(
        "Bachelor of Engineering in Computer EngineeringPune, India"
    )
    experience_heading = find("EXPERIENCE")
    experience_title = find("Artificial Intelligence InternFeb. 2026 – Mar. 2026")
    experience_company = find("Codec Technologies IndiaRemote")
    bullet = find(
        "Worked on Natural Language Processing (NLP) tasks as part of the internship project"
    )
    certifications_heading = find("CERTIFICATIONS")
    certification_bullet = find("Introduction to AI")
    projects_heading = find("PROJECTS")
    project_title = find("Smart Classroom Automation System")
    project_type = find("IoT-based Project")
    project_bullet = find(
        "Built an IoT-based system that detects motion in a classroom and automatically turns lights/appliances on or off accordingly"
    )
    skills_heading = find("TECHNICAL SKILLS")
    skill_line = find("Languages: Python, C++, SQL, HTML & CSS")

    # Preserve the original name and contact paragraph styles.
    name_paragraph = original[0]
    contact_paragraph = original[1]
    paragraph_text(name_paragraph, data["name"])
    paragraph_text(contact_paragraph, data["contact"])

    # Remove old body content, retaining the name and contact header.
    body = document._body._body
    for child in list(body):
        if child.tag == qn("w:p"):
            if child is name_paragraph._p or child is contact_paragraph._p:
                continue
            body.remove(child)
        elif child.tag != qn("w:sectPr"):
            body.remove(child)

    # Rebuild the sections in the template's original order.
    make_paragraph(document, heading, "EDUCATION")
    for item in data["education"]:
        make_paragraph(
            document, education_title,
            f'{item["institution"]}    {item["date"]}'
        )
        make_paragraph(
            document, education_detail,
            f'{item["degree"]}    {", ".join(item["details"])}'
            if item["details"] else item["degree"]
        )

    make_paragraph(document, experience_heading, "EXPERIENCE")
    for item in data["experience"]:
        make_paragraph(
            document, experience_title,
            f'{item["role"]}    {item["date"]}'
        )
        make_paragraph(document, experience_company, item["company"])
        for text in item["bullets"]:
            make_paragraph(document, bullet, text)

    make_paragraph(document, certifications_heading, "CERTIFICATIONS")
    for item in data["certifications"]:
        make_paragraph(document, certification_bullet, item)

    make_paragraph(document, projects_heading, "PROJECTS")
    for item in data["projects"]:
        make_paragraph(document, project_title, item["name"])
        if item["type"]:
            make_paragraph(document, project_type, item["type"])
        if item["description"]:
            make_paragraph(document, project_bullet, item["description"])
        if item["tags"]:
            make_paragraph(
                document, project_type,
                "Technologies: " + ", ".join(item["tags"])
            )

    make_paragraph(document, skills_heading, "TECHNICAL SKILLS")
    for label, values in data["skills"].items():
        make_paragraph(document, skill_line, f"{label}: {', '.join(values)}")

    if data["languages"]:
        make_paragraph(
            document, skill_line,
            f'Spoken Languages: {data["languages"]}'
        )

    
    # Apply Times New Roman to every run in the entire CV.
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)

            # Set the font for all Word script types.
            rpr = run._element.get_or_add_rPr()
            rfonts = rpr.rFonts

            if rfonts is None:
                from docx.oxml import OxmlElement
                rfonts = OxmlElement("w:rFonts")
                rpr.insert(0, rfonts)

            for font_attr in (
                "ascii",
                "hAnsi",
                "eastAsia",
                "cs",
            ):
                rfonts.set(qn(f"w:{font_attr}"), "Times New Roman")

    # Apply the font to text inside tables too.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(10)

                        rpr = run._element.get_or_add_rPr()
                        rfonts = rpr.rFonts

                        if rfonts is None:
                            from docx.oxml import OxmlElement
                            rfonts = OxmlElement("w:rFonts")
                            rpr.insert(0, rfonts)

                        for font_attr in (
                            "ascii",
                            "hAnsi",
                            "eastAsia",
                            "cs",
                        ):
                            rfonts.set(
                                qn(f"w:{font_attr}"),
                                "Times New Roman",
                            )

    document.save(OUTPUT_DOCX)

    # GitHub Actions installs LibreOffice to perform this conversion.
    subprocess.run([
        "libreoffice", "--headless",
        "--convert-to", "pdf",
        "--outdir", str(ROOT),
        str(OUTPUT_DOCX),
    ], check=True)

    converted = ROOT / f"{OUTPUT_DOCX.stem}.pdf"
    if not converted.exists():
        raise RuntimeError("LibreOffice did not generate the CV PDF.")

    converted.replace(OUTPUT_PDF)
    print(f"Generated {OUTPUT_PDF}")


if __name__ == "__main__":
    generate()
